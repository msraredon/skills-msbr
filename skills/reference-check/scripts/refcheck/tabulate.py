"""Build the two output tables.

Data model (matching how a reader thinks about it):
  * a **reference** is a unique numbered work (the bibliography number is its id);
  * a **citation** is one in-text marker tied to one sentence/claim, linking to
    one or more references;
  * **appropriateness** is judged per (citation, reference) pair.

So the **audit** sheet is a linear list of citations in document order, each
expanded to one row per referenced work. Citation-level cells (the claim, marker,
section) are merged across a citation's reference rows in the .xlsx, so you read
a claim once and scan its references and their appropriateness beneath it. The
**references** sheet is the de-duplicated library, one row per unique work.

The ``ref_summary`` / ``connection`` / ``appropriateness`` columns are filled by
the model; the deterministic layer computes existence, confidence, links, flags.
"""
from __future__ import annotations

import csv
from typing import Any, Optional

AUDIT_COLUMNS = [
    "citation", "marker", "section", "claim",              # citation-level (merged)
    "reference", "authors", "year", "title", "journal",    # reference-level
    "doi", "pmid", "url", "existence", "confidence",
    "ref_summary",                                          # model (per reference)
    "connection", "appropriateness",                       # model (per pair)
    "flag", "flag_reason", "notes",
]

# Columns merged across a citation's reference rows in the xlsx.
_MERGE_COLS = ["citation", "marker", "section", "claim"]

LIBRARY_COLUMNS = [
    "reference", "authors", "year", "title", "journal",
    "volume", "issue", "pages", "doi", "pmid", "url", "open_access",
    "existence", "confidence", "n_citations", "cited_by", "summary", "notes",
]

_LINK_COLS = {"doi", "pmid", "url", "open_access"}


# --------------------------------------------------------------------------- #
# field helpers
# --------------------------------------------------------------------------- #
def _authors_str(work: dict) -> str:
    names = []
    for a in work.get("author", []):
        if a.get("literal"):
            names.append(a["literal"])
        else:
            names.append(f"{a.get('family','')} {a.get('given','')}".strip())
    if not names and work.get("custom", {}).get("bib_raw"):
        # pre-resolution: show the leading authors from the raw string
        return work["custom"]["bib_raw"][:60] + "…"
    if len(names) > 8:
        names = names[:8] + ["et al."]
    return "; ".join(names)


def _year(work: dict) -> str:
    dp = work.get("issued", {}).get("date-parts")
    return str(dp[0][0]) if dp and dp[0] else ""


def _links(work: dict) -> dict:
    return (work.get("custom", {}).get("links") or {})


def _existence(work: Optional[dict]) -> str:
    if not work:
        return "unresolved (not in bibliography)"
    v = work.get("custom", {}).get("verification", {})
    status = v.get("status", "unverified")
    if status == "verified":
        tags = [t for t, on in (("DOI", v.get("doi_verified")),
                                ("PMID", v.get("pmid_verified"))) if on]
        return "verified (" + "+".join(tags) + ")" if tags else "verified"
    return status


def _confidence(work: Optional[dict]) -> str:
    if not work:
        return "n/a"
    custom = work.get("custom", {})
    v = custom.get("verification", {})
    src = custom.get("source", "")
    if v.get("status") != "verified":
        return "low (unverified)"
    if src == "endnote-embedded":
        return "high (embedded id, verified)"
    score = v.get("match_score")
    if score is not None:
        return f"medium (bib match {score}, verified)"
    return "high (verified)"


def _flag(work: Optional[dict]) -> tuple[str, str]:
    reasons = []
    if not work:
        return "TRUE", "reference number not found in bibliography"
    if not (work.get("DOI") or work.get("custom", {}).get("pmid")):
        reasons.append("no working identifier")
    v = work.get("custom", {}).get("verification", {})
    if v.get("status") != "verified":
        reasons.append("not verified online")
    score = v.get("match_score")
    if score is not None and score < 0.75:
        reasons.append(f"low bibliographic match ({score})")
    return ("TRUE" if reasons else "FALSE"), "; ".join(reasons)


# --------------------------------------------------------------------------- #
# row builders
# --------------------------------------------------------------------------- #
def build_audit_rows(citations: list, references: dict[int, dict],
                     judgments: Optional[dict] = None) -> list[dict[str, Any]]:
    """Build the per-citation audit rows.

    ``judgments`` (optional) carries the model's appropriateness output:
      {"references": {"16": {"summary": ...}},
       "pairs": {"C1|16": {"connection": ..., "appropriateness": ...}}}
    When present, it fills ref_summary/connection/appropriateness and escalates
    the flag for any pair rated below "Supports".
    """
    jrefs = (judgments or {}).get("references", {})
    jpairs = (judgments or {}).get("pairs", {})
    rows: list[dict[str, Any]] = []
    for c in citations:
        cd = c.to_dict() if hasattr(c, "to_dict") else c
        nums = cd.get("ref_numbers") or []
        cid = cd["id"]
        base = {
            "_group": cid,
            "citation": cid,
            "marker": cd["marker"],
            "section": _short_section(cd.get("section")),
            "claim": cd["sentence"],
        }
        if not nums:
            row = dict.fromkeys(AUDIT_COLUMNS, "")
            row.update(base)
            row["existence"] = "unresolved"
            row["flag"], row["flag_reason"] = "TRUE", cd.get("note") or "no reference link"
            rows.append(row)
            continue
        for n in nums:
            w = references.get(n)
            flag, reason = _flag(w)
            pj = jpairs.get(f"{cid}|{n}", {})
            appr = pj.get("appropriateness", "")
            reasons = [reason] if reason else []
            if appr and appr.lower() != "supports":
                flag = "TRUE"
                reasons.append(f"appropriateness: {appr}")
            row = dict(base)
            row.update({
                "reference": str(n),
                "authors": _authors_str(w) if w else "",
                "year": _year(w) if w else "",
                "title": (w.get("title", "") if w else ""),
                "journal": (w.get("container-title", "") if w else ""),
                "doi": (w.get("DOI", "") if w else ""),
                "pmid": (w.get("custom", {}).get("pmid", "") if w else ""),
                "url": _best_url(w),
                "existence": _existence(w),
                "confidence": _confidence(w),
                "ref_summary": jrefs.get(str(n), {}).get("summary", ""),
                "connection": pj.get("connection", ""),
                "appropriateness": appr,
                "flag": flag,
                "flag_reason": "; ".join(reasons),
                "notes": pj.get("notes", ""),
            })
            rows.append(row)
    return rows


def build_library_rows(references: dict[int, dict], citations: list,
                       judgments: Optional[dict] = None) -> list[dict[str, Any]]:
    jrefs = (judgments or {}).get("references", {})
    cited_by: dict[int, list[str]] = {}
    for c in citations:
        cd = c.to_dict() if hasattr(c, "to_dict") else c
        for n in cd.get("ref_numbers") or []:
            cited_by.setdefault(n, []).append(cd["id"])
    rows = []
    for n in sorted(references):
        if n < 0:
            continue
        w = references[n]
        rows.append({
            "reference": str(n),
            "authors": _authors_str(w),
            "year": _year(w),
            "title": w.get("title", ""),
            "journal": w.get("container-title", ""),
            "volume": w.get("volume", ""),
            "issue": w.get("issue", ""),
            "pages": w.get("page", ""),
            "doi": w.get("DOI", ""),
            "pmid": w.get("custom", {}).get("pmid", ""),
            "url": _best_url(w),
            "open_access": (w.get("custom", {}).get("links", {}) or {}).get("oa_url")
                           or (f"https://www.ncbi.nlm.nih.gov/pmc/articles/{w['custom']['pmcid']}/"
                               if w.get("custom", {}).get("pmcid") else ""),
            "existence": _existence(w),
            "confidence": _confidence(w),
            "n_citations": len(cited_by.get(n, [])),
            "cited_by": ", ".join(cited_by.get(n, [])),
            "summary": jrefs.get(str(n), {}).get("summary", ""),
            "notes": "",
        })
    return rows


def _best_url(work: Optional[dict]) -> str:
    if not work:
        return ""
    links = _links(work)
    return links.get("doi_url") or links.get("pubmed_url") or work.get("URL", "")


def _short_section(style: Optional[str]) -> str:
    if not style:
        return ""
    return style.replace("Heading", "H").replace("EndNote", "")


# --------------------------------------------------------------------------- #
# writers
# --------------------------------------------------------------------------- #
def write_csv(rows: list[dict], columns: list[str], path: str) -> None:
    with open(path, "w", newline="", encoding="utf-8-sig") as fh:
        w = csv.DictWriter(fh, fieldnames=columns, extrasaction="ignore")
        w.writeheader()
        for r in rows:
            w.writerow(r)


def write_xlsx(audit_rows, library_rows, path: str) -> bool:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
        from openpyxl.utils import get_column_letter
    except Exception:
        return False

    wb = Workbook()
    wb.remove(wb.active)
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(color="FFFFFF", bold=True)
    flag_fill = PatternFill("solid", fgColor="FCE4D6")
    band_fill = PatternFill("solid", fgColor="F2F6FC")
    link_font = Font(color="0563C1", underline="single")
    top = Alignment(vertical="top", wrap_text=True)
    thin_top = Border(top=Side(style="thin", color="BBBBBB"))

    def header(ws, columns):
        ws.append(columns)
        for j in range(1, len(columns) + 1):
            c = ws.cell(row=1, column=j)
            c.fill = header_fill
            c.font = header_font
            c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.freeze_panes = "A2"

    def hyperlink(cell, col, val):
        if col in _LINK_COLS and val:
            href = str(val)
            if col == "doi" and not href.startswith("http"):
                href = "https://doi.org/" + href
            elif col == "pmid" and not href.startswith("http"):
                href = f"https://pubmed.ncbi.nlm.nih.gov/{href}/"
            cell.hyperlink = href
            cell.font = link_font

    # ---- Audit sheet: grouped by citation, merged citation cells ----------
    ws = wb.create_sheet("Citations (audit)")
    header(ws, AUDIT_COLUMNS)
    col_idx = {c: i + 1 for i, c in enumerate(AUDIT_COLUMNS)}
    r = 2
    band = False
    i = 0
    while i < len(audit_rows):
        group = audit_rows[i]["_group"]
        j = i
        while j < len(audit_rows) and audit_rows[j]["_group"] == group:
            j += 1
        band = not band
        first_row = r
        for k in range(i, j):
            row = audit_rows[k]
            ws.append([row.get(c, "") for c in AUDIT_COLUMNS])
            for col, cidx in col_idx.items():
                cell = ws.cell(row=r, column=cidx)
                cell.alignment = top
                hyperlink(cell, col, row.get(col, ""))
                if band:
                    cell.fill = band_fill
                cell.border = thin_top
                if col == "flag" and str(row.get("flag")).upper() == "TRUE":
                    ws.cell(row=r, column=col_idx["appropriateness"]).fill = flag_fill
                    ws.cell(row=r, column=col_idx["existence"]).fill = flag_fill
            r += 1
        if j - i > 1:                       # merge citation-level cells
            for col in _MERGE_COLS:
                ci = col_idx[col]
                ws.merge_cells(start_row=first_row, start_column=ci,
                               end_row=r - 1, end_column=ci)
                ws.cell(row=first_row, column=ci).alignment = top
        i = j
    _widths(ws, AUDIT_COLUMNS, get_column_letter,
            {"claim": 55, "title": 40, "ref_summary": 45, "connection": 45,
             "authors": 26, "journal": 20, "flag_reason": 26, "notes": 20,
             "appropriateness": 14})

    # ---- References sheet -------------------------------------------------
    ws2 = wb.create_sheet("References (library)")
    header(ws2, LIBRARY_COLUMNS)
    ci2 = {c: i + 1 for i, c in enumerate(LIBRARY_COLUMNS)}
    for idx, row in enumerate(library_rows, start=2):
        ws2.append([row.get(c, "") for c in LIBRARY_COLUMNS])
        for col, cidx in ci2.items():
            cell = ws2.cell(row=idx, column=cidx)
            cell.alignment = top
            hyperlink(cell, col, row.get(col, ""))
            if str(row.get("flag", "")).upper() == "TRUE" or row.get("existence", "").startswith(("unresolved", "unverified")):
                ws2.cell(row=idx, column=ci2["existence"]).fill = flag_fill
    ws2.auto_filter.ref = ws2.dimensions
    _widths(ws2, LIBRARY_COLUMNS, get_column_letter,
            {"title": 45, "authors": 28, "journal": 22, "summary": 50,
             "cited_by": 20, "notes": 20})
    wb.save(path)
    return True


def _widths(ws, columns, get_col, widths):
    from openpyxl.styles import Alignment
    for j, col in enumerate(columns, 1):
        ws.column_dimensions[get_col(j)].width = widths.get(col, 12)
