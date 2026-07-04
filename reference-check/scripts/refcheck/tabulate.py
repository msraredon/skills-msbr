"""Build the two output tables from citations + resolved works.

  * audit table (per-citation): one row per (citation x cited work). This is the
    unit the appropriateness check reasons over — sentence vs. one reference.
    Unresolved citations still get a row so nothing is silently dropped.
  * library table (per-work): one row per unique work, the human-readable view
    of the reference library.

The ``summary``, ``connection``, and ``appropriateness`` columns are filled by
the model in the hybrid workflow (they require judgment over the abstract /
full text); the deterministic layer leaves them blank but always computes the
existence status, confidence, links, and flags.
"""
from __future__ import annotations

import csv
from typing import Any, Optional

AUDIT_COLUMNS = [
    "citation_id", "marker", "section", "cited_sentence",
    "authors", "title", "year", "journal",
    "doi", "pmid", "url",
    "existence_status", "match_confidence",
    "summary", "connection", "appropriateness",
    "flag", "flag_reason", "notes",
]

LIBRARY_COLUMNS = [
    "work_id", "authors", "title", "year", "journal",
    "volume", "issue", "pages", "doi", "pmid", "url",
    "existence_status", "n_citations", "summary", "notes",
]

_LINK_COLS = {"doi", "pmid", "url"}   # columns rendered as hyperlinks in xlsx


def _authors_str(work: dict) -> str:
    names = []
    for a in work.get("author", []):
        if a.get("literal"):
            names.append(a["literal"])
        else:
            names.append(f"{a.get('family','')} {a.get('given','')}".strip())
    if len(names) > 8:
        names = names[:8] + ["et al."]
    return "; ".join(names)


def _year(work: dict) -> str:
    dp = work.get("issued", {}).get("date-parts")
    return str(dp[0][0]) if dp and dp[0] else ""


def _links(work: dict) -> dict:
    return (work.get("custom", {}).get("links") or {})


def _existence(work: Optional[dict]) -> str:
    if not work:
        return "unresolved (no metadata)"
    v = work.get("custom", {}).get("verification", {})
    status = v.get("status", "unknown")
    if status == "verified":
        tags = []
        if v.get("doi_verified"):
            tags.append("DOI")
        if v.get("pmid_verified"):
            tags.append("PMID")
        return "verified (" + "+".join(tags) + ")"
    return status


def _confidence(work: Optional[dict]) -> str:
    if not work:
        return "n/a"
    src = work.get("custom", {}).get("source", "")
    v = work.get("custom", {}).get("verification", {})
    if src == "endnote-embedded" and v.get("status") == "verified":
        return "high (embedded id, verified online)"
    if v.get("status") == "verified":
        return "high (verified online)"
    return "low (unverified)"


def build_audit_rows(citations: list, works: dict[str, dict]) -> list[dict[str, Any]]:
    rows = []
    for c in citations:
        cd = c.to_dict() if hasattr(c, "to_dict") else c
        base = {
            "citation_id": cd["id"],
            "marker": cd["marker"],
            "section": cd.get("section") or "",
            "cited_sentence": cd["sentence"],
        }
        if not cd["work_ids"]:
            row = dict.fromkeys(AUDIT_COLUMNS, "")
            row.update(base)
            row["existence_status"] = "unresolved (no metadata)"
            row["match_confidence"] = "n/a"
            row["flag"] = "TRUE"
            row["flag_reason"] = "no embedded metadata; supply reference library or resolve online"
            row["notes"] = cd.get("note") or ""
            rows.append(row)
            continue
        for wid in cd["work_ids"]:
            w = works.get(wid, {})
            links = _links(w)
            existence = _existence(w)
            flag = "FALSE"
            reasons = []
            if not (w.get("DOI") or w.get("custom", {}).get("pmid")):
                flag = "TRUE"; reasons.append("no working identifier")
            if w.get("custom", {}).get("verification", {}).get("status") != "verified":
                flag = "TRUE"; reasons.append("identifier not verified online")
            row = {
                "citation_id": cd["id"],
                "marker": cd["marker"],
                "section": cd.get("section") or "",
                "cited_sentence": cd["sentence"],
                "authors": _authors_str(w),
                "title": w.get("title", ""),
                "year": _year(w),
                "journal": w.get("container-title", ""),
                "doi": w.get("DOI", ""),
                "pmid": w.get("custom", {}).get("pmid", ""),
                "url": links.get("doi_url") or links.get("pubmed_url") or w.get("URL", ""),
                "existence_status": existence,
                "match_confidence": _confidence(w),
                "summary": "",            # filled by model
                "connection": "",         # filled by model
                "appropriateness": "",    # filled by model
                "flag": flag,
                "flag_reason": "; ".join(reasons),
                "notes": "",
            }
            rows.append(row)
    return rows


def build_library_rows(works: dict[str, dict], citations: list) -> list[dict[str, Any]]:
    counts: dict[str, int] = {}
    for c in citations:
        cd = c.to_dict() if hasattr(c, "to_dict") else c
        for wid in cd["work_ids"]:
            counts[wid] = counts.get(wid, 0) + 1
    rows = []
    for wid, w in works.items():
        links = _links(w)
        rows.append({
            "work_id": wid,
            "authors": _authors_str(w),
            "title": w.get("title", ""),
            "year": _year(w),
            "journal": w.get("container-title", ""),
            "volume": w.get("volume", ""),
            "issue": w.get("issue", ""),
            "pages": w.get("page", ""),
            "doi": w.get("DOI", ""),
            "pmid": w.get("custom", {}).get("pmid", ""),
            "url": links.get("doi_url") or links.get("pubmed_url") or w.get("URL", ""),
            "existence_status": _existence(w),
            "n_citations": counts.get(wid, 0),
            "summary": "",
            "notes": "",
        })
    rows.sort(key=lambda r: (-r["n_citations"], r["authors"]))
    return rows


# --------------------------------------------------------------------------- #
# writers
# --------------------------------------------------------------------------- #
def write_csv(rows: list[dict], columns: list[str], path: str) -> None:
    with open(path, "w", newline="", encoding="utf-8-sig") as fh:
        w = csv.DictWriter(fh, fieldnames=columns, extrasaction="ignore")
        w.writeheader()
        for r in rows:
            w.writerow(r)


def write_xlsx(sheets: dict[str, tuple[list[dict], list[str]]], path: str) -> bool:
    """Write a multi-sheet workbook with hyperlinked id columns. Returns success."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except Exception:
        return False

    wb = Workbook()
    wb.remove(wb.active)
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(color="FFFFFF", bold=True)
    flag_fill = PatternFill("solid", fgColor="FCE4D6")
    link_font = Font(color="0563C1", underline="single")

    for name, (rows, columns) in sheets.items():
        ws = wb.create_sheet(name[:31])
        ws.append(columns)
        for j, _ in enumerate(columns, 1):
            c = ws.cell(row=1, column=j)
            c.fill = header_fill
            c.font = header_font
        for r in rows:
            ws.append([r.get(col, "") for col in columns])
            excel_row = ws.max_row
            for j, col in enumerate(columns, 1):
                cell = ws.cell(row=excel_row, column=j)
                val = r.get(col, "")
                if col in _LINK_COLS and val:
                    href = val
                    if col == "doi" and not str(val).startswith("http"):
                        href = "https://doi.org/" + str(val)
                    elif col == "pmid" and not str(val).startswith("http"):
                        href = f"https://pubmed.ncbi.nlm.nih.gov/{val}/"
                    cell.hyperlink = href
                    cell.font = link_font
                if col == "flag" and str(val).upper() == "TRUE":
                    for jj in range(1, len(columns) + 1):
                        ws.cell(row=excel_row, column=jj).fill = flag_fill
        # column widths + wrapping for text-heavy columns
        widths = {"cited_sentence": 60, "title": 45, "summary": 50,
                  "connection": 50, "authors": 30, "journal": 25,
                  "flag_reason": 30, "notes": 25}
        for j, col in enumerate(columns, 1):
            ws.column_dimensions[get_column_letter(j)].width = widths.get(col, 15)
            if col in widths and widths[col] >= 40:
                for row in ws.iter_rows(min_col=j, max_col=j, min_row=2):
                    row[0].alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
    wb.save(path)
    return True
